"""
Dispatch Planning ERP - Engineering Handover Document Generator
Generates both .docx (Word) and .pdf formats with complete file-by-file breakdown,
architecture details, priority audit files, and backend integration guide.
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether, HRFlowable
)
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.pdfgen import canvas


def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)


def create_word_document(filename="Dispatch_Planning_Engineering_Handover_Document.docx"):
    doc = docx.Document()

    # Page Margins (0.75 in)
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # Styles & Colors
    PRIMARY = RGBColor(30, 41, 59)      # Slate 800
    ACCENT = RGBColor(217, 119, 6)      # Amber 600
    SUBTEXT = RGBColor(71, 85, 105)     # Slate 600

    # Title
    title = doc.add_paragraph()
    r_title = title.add_run("DISPATCH PLANNING ERP")
    r_title.bold = True
    r_title.font.size = Pt(24)
    r_title.font.color.rgb = PRIMARY
    r_title.font.name = "Arial"
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_after = Pt(2)

    subtitle = doc.add_paragraph()
    r_sub = subtitle.add_run("Senior Engineering Handover & Technical Blueprint")
    r_sub.font.size = Pt(14)
    r_sub.font.color.rgb = ACCENT
    r_sub.bold = True
    subtitle.paragraph_format.space_after = Pt(8)

    meta = doc.add_paragraph()
    r_meta = meta.add_run("Audience: Frontend & Backend Engineering Teams | Status: Production-Ready (V6 Release)")
    r_meta.font.size = Pt(9.5)
    r_meta.font.color.rgb = SUBTEXT
    r_meta.italic = True
    meta.paragraph_format.space_after = Pt(18)

    # Section 1: Executive Summary
    h1 = doc.add_heading("1. Executive Summary & Tech Stack", level=1)
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(6)

    p1 = doc.add_paragraph(
        "This project is an Enterprise-Grade Dispatch Planning and Logistics ERP System built for managing supply chain workflows, "
        "vehicle fleet scheduling, multi-drop freight allocations, Proforma Invoice (PI) approvals, gate passes, incident rollbacks, "
        "PO Pickups, and Special Supply Chain Requests."
    )
    p1.paragraph_format.space_after = Pt(8)

    # Tech Stack Table
    table_stack = doc.add_table(rows=1, cols=3)
    table_stack.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_stack.autofit = False

    hdr_cells = table_stack.rows[0].cells
    hdr_cells[0].text = "Layer"
    hdr_cells[1].text = "Technology / Tool"
    hdr_cells[2].text = "Key Purpose & Responsibility"
    for cell in hdr_cells:
        set_cell_background(cell, "1E293B")
        for p in cell.paragraphs:
            p.runs[0].font.color.rgb = RGBColor(255, 255, 255)
            p.runs[0].font.bold = True
            p.runs[0].font.size = Pt(9.5)

    stack_data = [
        ("Frontend Core", "React 19, TypeScript 5.8, Vite 6", "Modern type-safe SPA with fast HMR and optimized builds"),
        ("Styling & UI", "Tailwind CSS v4, Lucide React Icons", "Utility-first responsive styling with enterprise color scheme"),
        ("State & Mock API", "Singleton Pub/Sub + LocalStorage (api.ts)", "In-memory mock API ready for 1-to-1 REST endpoint replacement"),
        ("Routing", "Hash / Popstate Custom SPA Router (App.tsx)", "Zero external routing dependency; fully compatible with GitHub Pages"),
        ("Deployment", "GitHub Pages (gh-pages branch)", "Automated production builds with SPA 404.html fallback routing")
    ]

    for layer, tech, purp in stack_data:
        row = table_stack.add_row().cells
        row[0].text = layer
        row[1].text = tech
        row[2].text = purp
        for cell in row:
            set_cell_margins(cell, 80, 80, 100, 100)
            for p in cell.paragraphs:
                p.runs[0].font.size = Pt(9)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Section 2: Core Architecture & Data Flow
    h2 = doc.add_heading("2. Architecture & Data Flow (How It Works)", level=1)
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(6)

    p2 = doc.add_paragraph(
        "1. Single Data Service Pattern (api.ts): All operations go through `dispatchService` (Singleton in `src/services/api.ts`). "
        "It maintains in-memory state initialized from LocalStorage (or `mockData.ts` if empty). When any state change occurs, "
        "it persists to LocalStorage and triggers subscribers (pub/sub) so all active views update simultaneously without reload.\n"
        "2. Centralized Modal System (AppLayout.tsx): All modal states (Create Dispatch, New PI, Gate Pass, Rollback, Special Request) "
        "are owned by `AppLayout.tsx` and exposed via helper callbacks (`helpers.openNewDispatch`, etc.) to child pages.\n"
        "3. Route Architecture (App.tsx): Routing uses `window.location.hash`. The `renderPage(helpers)` function acts as the root switchboard."
    )
    p2.paragraph_format.space_after = Pt(12)

    # Section 3: Most Important Files to Inspect First
    h3 = doc.add_heading("3. Priority Files to Inspect First (Quick Onboarding)", level=1)
    h3.paragraph_format.space_before = Pt(12)
    h3.paragraph_format.space_after = Pt(6)

    doc.add_paragraph(
        "If a new engineer (Frontend or Backend) joins today, they should read these 5 files in exact order:"
    )

    t_prio = doc.add_table(rows=1, cols=4)
    t_prio.alignment = WD_TABLE_ALIGNMENT.CENTER
    h_prio = t_prio.rows[0].cells
    h_prio[0].text = "Priority"
    h_prio[1].text = "File Path"
    h_prio[2].text = "Why It Matters"
    h_prio[3].text = "What to Check / Modify"
    for cell in h_prio:
        set_cell_background(cell, "D97706")
        for p in cell.paragraphs:
            p.runs[0].font.color.rgb = RGBColor(255, 255, 255)
            p.runs[0].font.bold = True
            p.runs[0].font.size = Pt(9.5)

    prio_data = [
        ("1 (Critical)", "src/types/index.ts", "Central data contracts for the entire ERP system.", "Review all schemas: ProformaInvoice, DispatchPlan, Vehicle, POPickup, SpecialRequest."),
        ("2 (Critical)", "src/services/api.ts", "Business logic, CRUD operations, state persistence, mock endpoints.", "Backend team will replace these mock methods with real Axios / Fetch API endpoints."),
        ("3 (High)", "src/layouts/AppLayout.tsx & src/App.tsx", "Application shell, route switchboard, global modal registry.", "Where new pages, navigation paths, and modals are registered."),
        ("4 (High)", "src/components/modals/CreateDispatchModal.tsx", "Core dispatch engine with vehicle capacity & weight calculation.", "Inspect vehicle payload validation, multi-PI multi-drop selection logic."),
        ("5 (Medium)", "src/pages/DashboardPage.tsx", "Main executive control tower with real-time KPI metrics.", "KPI cards, active trip queue, pending approval lists, quick action triggers.")
    ]

    for p_num, f_path, w_matters, w_check in prio_data:
        row = t_prio.add_row().cells
        row[0].text = p_num
        row[1].text = f_path
        row[2].text = w_matters
        row[3].text = w_check
        for cell in row:
            set_cell_margins(cell, 80, 80, 100, 100)
            for p in cell.paragraphs:
                p.runs[0].font.size = Pt(8.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Section 4: Complete Step-by-Step File Inventory
    h4 = doc.add_heading("4. Comprehensive File Inventory & Purpose Breakdown", level=1)
    h4.paragraph_format.space_before = Pt(12)
    h4.paragraph_format.space_after = Pt(6)

    file_categories = [
        ("A. Core Engine & State Layer", [
            ("src/types/index.ts", "Contains all TypeScript interfaces, enum values, and data models (DispatchPlan, ProformaInvoice, Vehicle, Driver, Client, Warehouse, RollbackLog, POPickup, SpecialRequest)."),
            ("src/services/api.ts", "The Central DispatchDataService class. Manages all CRUD actions, state reactivity via pub/sub listeners, LocalStorage read/write, and helper status updaters."),
            ("src/data/mockData.ts", "Initial seed data for Vehicles, Drivers, Clients, PIs, Dispatches, POPickups, and Special Requests when LocalStorage is empty.")
        ]),
        ("B. Shell, Layout & Routing", [
            ("src/App.tsx", "Root React component. Handles hash-based routing, URL synchronization, and renders active page with modal helper bindings."),
            ("src/layouts/AppLayout.tsx", "Master application wrapper containing Sidebar, Navbar, and all global modal dialogs. Provides helper callbacks to children."),
            ("src/components/Navbar.tsx", "Top navigation bar showing Hub badge (Bhiwandi Hub), notification drawer, quick action buttons, and user profile."),
            ("src/components/Sidebar.tsx", "Left sidebar with 5 primary nav items + collapsible 'More Operations' dropdown (Special_Request, PO-Pickup, Export Transfer, KB Orders).")
        ]),
        ("C. Operational Pages (src/pages/)", [
            ("DashboardPage.tsx", "Executive command center. Features 6 KPI cards (Pending PI, High Priority, Pending PO, Dispatches, Fleet Available, Special Requests), active loading queue, and pending approvals."),
            ("DispatchPlanningPage.tsx", "Primary trip planning page. Interactive vehicle selector, capacity gauge, PI checklist, and multi-PI trip creation."),
            ("PendingPIPage.tsx", "Registry of pending Proforma Invoices with Priority filter, approval workflow modal, and batch dispatch planning."),
            ("SpecialRequestPage.tsx", "Web form registry for supply chain work orders. Displays Vendor Name, Location, Contact, Items, Deadline, Feedback, Attachments, and 'Mark Work Done' button."),
            ("POPickupPage.tsx", "Purchase Order pickup registry. Features Planned vs Actual date tracking with automatic PENDING / COMPLETED status updates."),
            ("TodaysPlanningPage.tsx", "Real-time dispatch monitor for dispatches scheduled today. View gate pass, trigger rollback, update trip status."),
            ("UpcomingPlanningPage.tsx", "Forward-looking trip calendar for upcoming dispatches."),
            ("CompletedDispatchPage.tsx", "Registry of completed deliveries with POD (Proof of Delivery) tracking."),
            ("DispatchHistoryPage.tsx", "Historical searchable log of all dispatched trips."),
            ("RollbackPage.tsx", "Audit trail of rolled-back / cancelled dispatch plans with mandatory rollback reason tracking."),
            ("ReportsPage.tsx", "Fleet efficiency metrics, driver trip counts, volume statistics."),
            ("SettingsPage.tsx", "Master data configuration: Fleet vehicles, drivers, clients, and warehouse management."),
            ("ExportTransferPage.tsx & KBOrdersPage.tsx", "Specialized logistics views for export transfers and KB order handling.")
        ]),
        ("D. Interactive Modals (src/components/modals/)", [
            ("CreateDispatchModal.tsx", "Multi-PI dispatch planning modal with vehicle auto-calculation, weight/volume validation, multi-drop stop sequence ordering."),
            ("NewPIModal.tsx", "Invoice creation form with multi-line item entry, GST calculation, destination warehouse, and priority selection."),
            ("SpecialRequestModal.tsx", "Dedicated form for supply chain tasks: Vendor name, GPS location, contact info, item details, deadline date, feedback, and file attachments."),
            ("GatePassModal.tsx", "Official security gate pass with vehicle barcode, driver details, seal number, and print-ready format."),
            ("RollbackModal.tsx", "Trip rollback dialog enforcing audit reasoning before resetting allocated PIs back to pending status."),
            ("ApprovePIModal.tsx", "Financial & logistics approval dialog for pending invoices."),
            ("ViewPIDetailsModal.tsx", "Read-only detailed view of invoice line items, delivery location, and status history.")
        ]),
        ("E. Common UI Widgets (src/components/common/)", [
            ("CapacityBar.tsx", "Progress bar showing weight and volume capacity utilization percentages (Green < 80%, Yellow < 95%, Red >= 95%)."),
            ("StatusBadge.tsx", "Color-coded status badge supporting all ERP statuses (PENDING, APPROVED, LOADING, DISPATCHED, COMPLETED, ROLLED_BACK, etc.)."),
            ("StatCard.tsx", "Reusable KPI stat metric card component.")
        ])
    ]

    for cat_title, files_list in file_categories:
        h_cat = doc.add_heading(cat_title, level=2)
        h_cat.paragraph_format.space_before = Pt(8)
        h_cat.paragraph_format.space_after = Pt(4)

        t_cat = doc.add_table(rows=1, cols=2)
        t_cat.alignment = WD_TABLE_ALIGNMENT.CENTER
        h_c = t_cat.rows[0].cells
        h_c[0].text = "File Name"
        h_c[1].text = "Detailed Working & Responsibility"
        set_cell_background(h_c[0], "334155")
        set_cell_background(h_c[1], "334155")
        for cell in h_c:
            for p in cell.paragraphs:
                p.runs[0].font.color.rgb = RGBColor(255, 255, 255)
                p.runs[0].font.bold = True
                p.runs[0].font.size = Pt(9)

        for fn, desc in files_list:
            row = t_cat.add_row().cells
            row[0].text = fn
            row[1].text = desc
            for cell in row:
                set_cell_margins(cell, 60, 60, 90, 90)
                for p in cell.paragraphs:
                    p.runs[0].font.size = Pt(8.5)

        doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Section 5: Audit & Code Customization Guide
    h5 = doc.add_heading("5. Code Audit & Customization Guide (What is Safe to Modify/Remove)", level=1)
    h5.paragraph_format.space_before = Pt(12)
    h5.paragraph_format.space_after = Pt(6)

    t_audit = doc.add_table(rows=1, cols=3)
    t_audit.alignment = WD_TABLE_ALIGNMENT.CENTER
    h_a = t_audit.rows[0].cells
    h_a[0].text = "Category"
    h_a[1].text = "Target Files"
    h_a[2].text = "Engineering Guidelines"
    for cell in h_a:
        set_cell_background(cell, "0F766E")
        for p in cell.paragraphs:
            p.runs[0].font.color.rgb = RGBColor(255, 255, 255)
            p.runs[0].font.bold = True
            p.runs[0].font.size = Pt(9.5)

    audit_data = [
        ("Backend API Migration", "src/services/api.ts", "Replace LocalStorage methods with async HTTP calls (Axios/Fetch). Frontend UI components will continue functioning seamlessly if method signatures are maintained."),
        ("Safe to Remove / Customize", "src/data/mockData.ts\nsrc/pages/ExportTransferPage.tsx\nsrc/pages/KBOrdersPage.tsx", "Mock data can be replaced with API responses. ExportTransfer & KBOrders are template pages ready for client-specific business rules or removal."),
        ("DO NOT Modify Without Review", "src/components/modals/CreateDispatchModal.tsx\nsrc/types/index.ts", "Vehicle payload calculation, stop sequencing, and PI allocation algorithms are tightly coupled with the dispatch workflow."),
        ("Hub & Location Config", "src/components/Navbar.tsx\nsrc/pages/DashboardPage.tsx", "Hub name ('Bhiwandi Hub') is currently hardcoded in header badges. Can be made dynamic from user profile/auth context.")
    ]

    for cat, fls, gd in audit_data:
        row = t_audit.add_row().cells
        row[0].text = cat
        row[1].text = fls
        row[2].text = gd
        for cell in row:
            set_cell_margins(cell, 80, 80, 100, 100)
            for p in cell.paragraphs:
                p.runs[0].font.size = Pt(8.5)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Section 6: Backend Team Integration Blueprint
    h6 = doc.add_heading("6. Backend Team Integration Blueprint (API Contracts)", level=1)
    h6.paragraph_format.space_before = Pt(12)
    h6.paragraph_format.space_after = Pt(6)

    doc.add_paragraph(
        "To connect the Frontend to a live Backend (Node.js/Spring Boot/Python/Go), implement the following REST endpoints:"
    )

    t_api = doc.add_table(rows=1, cols=4)
    t_api.alignment = WD_TABLE_ALIGNMENT.CENTER
    h_api = t_api.rows[0].cells
    h_api[0].text = "Method"
    h_api[1].text = "Endpoint Route"
    h_api[2].text = "Frontend Consumer Method"
    h_api[3].text = "Payload / Response Details"
    for cell in h_api:
        set_cell_background(cell, "1E293B")
        for p in cell.paragraphs:
            p.runs[0].font.color.rgb = RGBColor(255, 255, 255)
            p.runs[0].font.bold = True
            p.runs[0].font.size = Pt(9)

    api_endpoints = [
        ("GET", "/api/v1/pis", "dispatchService.getPIs()", "Returns list of Proforma Invoices"),
        ("POST", "/api/v1/pis", "dispatchService.addPI(piData)", "Creates new PI with line items"),
        ("PATCH", "/api/v1/pis/:id/status", "dispatchService.updatePIStatus(id, status)", "Approve, reject, or mark dispatched"),
        ("GET", "/api/v1/dispatches", "dispatchService.getDispatchPlans()", "Returns all dispatch plans with stops & vehicles"),
        ("POST", "/api/v1/dispatches", "dispatchService.createDispatchPlan(plan)", "Creates trip, reserves vehicle, marks PIs planned"),
        ("POST", "/api/v1/dispatches/:id/rollback", "dispatchService.rollbackDispatch(id, reason)", "Rolls back trip, releases vehicle, resets PIs to PENDING"),
        ("GET", "/api/v1/fleet/vehicles", "dispatchService.getVehicles()", "Returns vehicle fleet with capacity & availability"),
        ("GET", "/api/v1/fleet/drivers", "dispatchService.getDrivers()", "Returns driver master data"),
        ("GET", "/api/v1/special-requests", "dispatchService.getSpecialRequests()", "Returns supply chain special work requests"),
        ("POST", "/api/v1/special-requests", "dispatchService.addSpecialRequest(req)", "Creates new special request with attachments"),
        ("PATCH", "/api/v1/special-requests/:id/done", "dispatchService.updateSpecialRequestStatus(id, 'WORK_DONE')", "Marks request as work completed")
    ]

    for m, r, f, p in api_endpoints:
        row = t_api.add_row().cells
        row[0].text = m
        row[1].text = r
        row[2].text = f
        row[3].text = p
        for cell in row:
            set_cell_margins(cell, 60, 60, 80, 80)
            for par in cell.paragraphs:
                par.runs[0].font.size = Pt(8)

    doc.save(filename)
    print(f"[SUCCESS] Word document generated: {filename}")


# ---------------- PDF GENERATOR WITH REPORTLAB ----------------

class NumberedCanvas(canvas.Canvas):
    def __init__(self, *args, **kwargs):
        super(NumberedCanvas, self).__init__(*args, **kwargs)
        self._saved_page_states = []

    def showPage(self):
        self._saved_page_states.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        num_pages = len(self._saved_page_states)
        for state in self._saved_page_states:
            self.__dict__.update(state)
            self.draw_page_number(num_pages)
            super(NumberedCanvas, self).showPage()
        super(NumberedCanvas, self).save()

    def draw_page_number(self, page_count):
        self.saveState()
        self.setFont("Helvetica", 8)
        self.setFillColor(colors.HexColor("#64748B"))
        # Header
        self.drawString(54, 750, "DISPATCH PLANNING ERP — Senior Engineering Handover Document")
        self.setStrokeColor(colors.HexColor("#CBD5E1"))
        self.setLineWidth(0.5)
        self.line(54, 742, 558, 742)

        # Footer
        page_text = f"Page {self._pageNumber} of {page_count}"
        self.drawRightString(558, 36, page_text)
        self.drawString(54, 36, "Confidential — Engineering Blueprint & Handover Spec")
        self.line(54, 48, 558, 48)
        self.restoreState()


def create_pdf_document(filename="Dispatch_Planning_Engineering_Handover_Document.pdf"):
    doc = SimpleDocTemplate(
        filename,
        pagesize=letter,
        leftMargin=54,
        rightMargin=54,
        topMargin=60,
        bottomMargin=60
    )

    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=20,
        leading=24,
        textColor=colors.HexColor('#0F172A')
    )

    subtitle_style = ParagraphStyle(
        'DocSubtitle',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=12,
        leading=16,
        textColor=colors.HexColor('#D97706')
    )

    meta_style = ParagraphStyle(
        'DocMeta',
        parent=styles['Normal'],
        fontName='Helvetica-Oblique',
        fontSize=8.5,
        leading=12,
        textColor=colors.HexColor('#64748B')
    )

    h1_style = ParagraphStyle(
        'SectionH1',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=12,
        leading=16,
        textColor=colors.HexColor('#1E293B'),
        spaceBefore=12,
        spaceAfter=6
    )

    h2_style = ParagraphStyle(
        'SectionH2',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=10,
        leading=14,
        textColor=colors.HexColor('#334155'),
        spaceBefore=8,
        spaceAfter=4
    )

    body_style = ParagraphStyle(
        'DocBody',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8.5,
        leading=12,
        textColor=colors.HexColor('#334155')
    )

    cell_hdr_style = ParagraphStyle(
        'CellHdr',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8,
        leading=10,
        textColor=colors.white
    )

    cell_body_style = ParagraphStyle(
        'CellBody',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=7.5,
        leading=10,
        textColor=colors.HexColor('#1E293B')
    )

    cell_body_bold = ParagraphStyle(
        'CellBodyBold',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=7.5,
        leading=10,
        textColor=colors.HexColor('#0F172A')
    )

    elements = []

    # Title Banner
    elements.append(Paragraph("DISPATCH PLANNING ERP", title_style))
    elements.append(Paragraph("Senior Engineering Handover & Technical Blueprint", subtitle_style))
    elements.append(Paragraph("Audience: Frontend & Backend Engineering Teams | Status: Production-Ready (V6 Release)", meta_style))
    elements.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor('#CBD5E1'), spaceAfter=10, spaceBefore=4))

    # Section 1
    elements.append(Paragraph("1. Executive Summary & Tech Stack", h1_style))
    elements.append(Paragraph(
        "This project is an Enterprise-Grade Dispatch Planning and Logistics ERP System built for managing supply chain workflows, "
        "vehicle fleet scheduling, multi-drop freight allocations, Proforma Invoice (PI) approvals, gate passes, incident rollbacks, "
        "PO Pickups, and Special Supply Chain Requests.",
        body_style
    ))
    elements.append(Spacer(1, 6))

    stack_rows = [
        [Paragraph("Layer", cell_hdr_style), Paragraph("Technology / Tool", cell_hdr_style), Paragraph("Key Purpose & Responsibility", cell_hdr_style)],
        [Paragraph("Frontend Core", cell_body_bold), Paragraph("React 19, TypeScript 5.8, Vite 6", cell_body_style), Paragraph("Modern type-safe SPA with fast HMR and optimized builds", cell_body_style)],
        [Paragraph("Styling & UI", cell_body_bold), Paragraph("Tailwind CSS v4, Lucide React Icons", cell_body_style), Paragraph("Utility-first responsive styling with enterprise color scheme", cell_body_style)],
        [Paragraph("State & Mock API", cell_body_bold), Paragraph("Singleton Pub/Sub + LocalStorage (api.ts)", cell_body_style), Paragraph("In-memory mock API ready for 1-to-1 REST endpoint replacement", cell_body_style)],
        [Paragraph("Routing", cell_body_bold), Paragraph("Hash / Popstate Custom SPA Router (App.tsx)", cell_body_style), Paragraph("Zero external routing dependency; fully compatible with GitHub Pages", cell_body_style)],
        [Paragraph("Deployment", cell_body_bold), Paragraph("GitHub Pages (gh-pages branch)", cell_body_style), Paragraph("Automated production builds with SPA 404.html fallback routing", cell_body_style)]
    ]

    t_stack = Table(stack_rows, colWidths=[90, 160, 254])
    t_stack.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E293B')),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
        ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('LEFTPADDING', (0, 0), (-1, -1), 6),
        ('RIGHTPADDING', (0, 0), (-1, -1), 6),
    ]))
    elements.append(t_stack)
    elements.append(Spacer(1, 10))

    # Section 2
    elements.append(Paragraph("2. Architecture & Data Flow (How It Works)", h1_style))
    elements.append(Paragraph(
        "<b>1. Single Data Service Pattern (api.ts):</b> All operations go through <code>dispatchService</code> (Singleton in <code>src/services/api.ts</code>). "
        "It maintains in-memory state initialized from LocalStorage (or <code>mockData.ts</code> if empty). When any state change occurs, "
        "it persists to LocalStorage and triggers subscribers (pub/sub) so all active views update simultaneously without reload.<br/>"
        "<b>2. Centralized Modal System (AppLayout.tsx):</b> All modal states (Create Dispatch, New PI, Gate Pass, Rollback, Special Request) "
        "are owned by <code>AppLayout.tsx</code> and exposed via helper callbacks (<code>helpers.openNewDispatch</code>, etc.) to child pages.<br/>"
        "<b>3. Route Architecture (App.tsx):</b> Routing uses <code>window.location.hash</code>. The <code>renderPage(helpers)</code> function acts as the root switchboard.",
        body_style
    ))
    elements.append(Spacer(1, 10))

    # Section 3: Priority Files
    elements.append(Paragraph("3. Priority Files to Inspect First (Quick Onboarding)", h1_style))
    elements.append(Paragraph("If a new engineer (Frontend or Backend) joins today, they should read these 5 files in exact order:", body_style))
    elements.append(Spacer(1, 6))

    prio_rows = [
        [Paragraph("Priority", cell_hdr_style), Paragraph("File Path", cell_hdr_style), Paragraph("Why It Matters", cell_hdr_style), Paragraph("What to Check / Modify", cell_hdr_style)],
        [Paragraph("1 (Critical)", cell_body_bold), Paragraph("src/types/index.ts", cell_body_bold), Paragraph("Central data contracts for the entire ERP system.", cell_body_style), Paragraph("Review all schemas: ProformaInvoice, DispatchPlan, Vehicle, POPickup, SpecialRequest.", cell_body_style)],
        [Paragraph("2 (Critical)", cell_body_bold), Paragraph("src/services/api.ts", cell_body_bold), Paragraph("Business logic, CRUD operations, state persistence, mock endpoints.", cell_body_style), Paragraph("Backend team will replace these mock methods with real Axios / Fetch API endpoints.", cell_body_style)],
        [Paragraph("3 (High)", cell_body_bold), Paragraph("src/layouts/AppLayout.tsx<br/>& src/App.tsx", cell_body_bold), Paragraph("Application shell, route switchboard, global modal registry.", cell_body_style), Paragraph("Where new pages, navigation paths, and modals are registered.", cell_body_style)],
        [Paragraph("4 (High)", cell_body_bold), Paragraph("src/components/modals/<br/>CreateDispatchModal.tsx", cell_body_bold), Paragraph("Core dispatch engine with vehicle capacity & weight calculation.", cell_body_style), Paragraph("Inspect vehicle payload validation, multi-PI multi-drop selection logic.", cell_body_style)],
        [Paragraph("5 (Medium)", cell_body_bold), Paragraph("src/pages/DashboardPage.tsx", cell_body_bold), Paragraph("Main executive control tower with real-time KPI metrics.", cell_body_style), Paragraph("KPI cards, active trip queue, pending approval lists, quick action triggers.", cell_body_style)]
    ]

    t_prio_pdf = Table(prio_rows, colWidths=[65, 115, 140, 184])
    t_prio_pdf.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#D97706')),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
        ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('LEFTPADDING', (0, 0), (-1, -1), 6),
        ('RIGHTPADDING', (0, 0), (-1, -1), 6),
    ]))
    elements.append(t_prio_pdf)
    elements.append(Spacer(1, 10))

    # Section 4: File Inventory
    elements.append(Paragraph("4. Comprehensive File Inventory & Purpose Breakdown", h1_style))

    categories_pdf = [
        ("A. Core Engine & State Layer", [
            ("src/types/index.ts", "Contains all TypeScript interfaces, enum values, and data models (DispatchPlan, ProformaInvoice, Vehicle, Driver, Client, Warehouse, RollbackLog, POPickup, SpecialRequest)."),
            ("src/services/api.ts", "The Central DispatchDataService class. Manages all CRUD actions, state reactivity via pub/sub listeners, LocalStorage read/write, and helper status updaters."),
            ("src/data/mockData.ts", "Initial seed data for Vehicles, Drivers, Clients, PIs, Dispatches, POPickups, and Special Requests when LocalStorage is empty.")
        ]),
        ("B. Shell, Layout & Navigation", [
            ("src/App.tsx", "Root React component. Handles hash-based routing, URL synchronization, and renders active page with modal helper bindings."),
            ("src/layouts/AppLayout.tsx", "Master application wrapper containing Sidebar, Navbar, and all global modal dialogs. Provides helper callbacks to children."),
            ("src/components/Navbar.tsx", "Top navigation bar showing Hub badge (Bhiwandi Hub), notification drawer, quick action buttons, and user profile."),
            ("src/components/Sidebar.tsx", "Left sidebar with 5 primary nav items + collapsible 'More Operations' dropdown (Special_Request, PO-Pickup, Export Transfer, KB Orders).")
        ]),
        ("C. Operational Pages (src/pages/)", [
            ("DashboardPage.tsx", "Executive command center. Features 6 KPI cards (Pending PI, High Priority, Pending PO, Dispatches, Fleet Available, Special Requests), active loading queue, and pending approvals."),
            ("DispatchPlanningPage.tsx", "Primary trip planning page. Interactive vehicle selector, capacity gauge, PI checklist, and multi-PI trip creation."),
            ("PendingPIPage.tsx", "Registry of pending Proforma Invoices with Priority filter, approval workflow modal, and batch dispatch planning."),
            ("SpecialRequestPage.tsx", "Web form registry for supply chain work orders. Displays Vendor Name, Location, Contact, Items, Deadline, Feedback, Attachments, and 'Mark Work Done' button."),
            ("POPickupPage.tsx", "Purchase Order pickup registry. Features Planned vs Actual date tracking with automatic PENDING / COMPLETED status updates."),
            ("TodaysPlanningPage.tsx", "Real-time dispatch monitor for dispatches scheduled today. View gate pass, trigger rollback, update trip status."),
            ("UpcomingPlanningPage.tsx", "Forward-looking trip calendar for upcoming dispatches."),
            ("CompletedDispatchPage.tsx", "Registry of completed deliveries with POD (Proof of Delivery) tracking."),
            ("DispatchHistoryPage.tsx", "Historical searchable log of all dispatched trips."),
            ("RollbackPage.tsx", "Audit trail of rolled-back / cancelled dispatch plans with mandatory rollback reason tracking."),
            ("ReportsPage.tsx", "Fleet efficiency metrics, driver trip counts, volume statistics."),
            ("SettingsPage.tsx", "Master data configuration: Fleet vehicles, drivers, clients, and warehouse management."),
            ("ExportTransferPage.tsx & KBOrdersPage.tsx", "Specialized logistics views for export transfers and KB order handling.")
        ]),
        ("D. Interactive Modals (src/components/modals/)", [
            ("CreateDispatchModal.tsx", "Multi-PI dispatch planning modal with vehicle auto-calculation, weight/volume validation, multi-drop stop sequence ordering."),
            ("NewPIModal.tsx", "Invoice creation form with multi-line item entry, GST calculation, destination warehouse, and priority selection."),
            ("SpecialRequestModal.tsx", "Dedicated form for supply chain tasks: Vendor name, GPS location, contact info, item details, deadline date, feedback, and file attachments."),
            ("GatePassModal.tsx", "Official security gate pass with vehicle barcode, driver details, seal number, and print-ready format."),
            ("RollbackModal.tsx", "Trip rollback dialog enforcing audit reasoning before resetting allocated PIs back to pending status."),
            ("ApprovePIModal.tsx", "Financial & logistics approval dialog for pending invoices."),
            ("ViewPIDetailsModal.tsx", "Read-only detailed view of invoice line items, delivery location, and status history.")
        ]),
        ("E. Common UI Widgets (src/components/common/)", [
            ("CapacityBar.tsx", "Progress bar showing weight and volume capacity utilization percentages (Green < 80%, Yellow < 95%, Red >= 95%)."),
            ("StatusBadge.tsx", "Color-coded status badge supporting all ERP statuses (PENDING, APPROVED, LOADING, DISPATCHED, COMPLETED, ROLLED_BACK, etc.)."),
            ("StatCard.tsx", "Reusable KPI stat metric card component.")
        ])
    ]

    for cat_t, f_list in categories_pdf:
        elements.append(Paragraph(cat_t, h2_style))
        cat_rows = [[Paragraph("File Name", cell_hdr_style), Paragraph("Detailed Working & Responsibility", cell_hdr_style)]]
        for fn, desc in f_list:
            cat_rows.append([Paragraph(fn, cell_body_bold), Paragraph(desc, cell_body_style)])

        t_cat_pdf = Table(cat_rows, colWidths=[150, 354])
        t_cat_pdf.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#334155')),
            ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
            ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
            ('TOPPADDING', (0, 0), (-1, -1), 3),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
            ('LEFTPADDING', (0, 0), (-1, -1), 5),
            ('RIGHTPADDING', (0, 0), (-1, -1), 5),
        ]))
        elements.append(t_cat_pdf)
        elements.append(Spacer(1, 6))

    elements.append(Spacer(1, 6))

    # Section 5: Audit & Customization
    elements.append(Paragraph("5. Code Audit & Customization Guide (What is Safe to Modify/Remove)", h1_style))

    audit_rows = [
        [Paragraph("Category", cell_hdr_style), Paragraph("Target Files", cell_hdr_style), Paragraph("Engineering Guidelines", cell_hdr_style)],
        [Paragraph("Backend API Migration", cell_body_bold), Paragraph("src/services/api.ts", cell_body_bold), Paragraph("Replace LocalStorage methods with async HTTP calls (Axios/Fetch). Frontend UI components will continue functioning seamlessly if method signatures are maintained.", cell_body_style)],
        [Paragraph("Safe to Remove / Customize", cell_body_bold), Paragraph("src/data/mockData.ts<br/>src/pages/ExportTransferPage.tsx<br/>src/pages/KBOrdersPage.tsx", cell_body_bold), Paragraph("Mock data can be replaced with API responses. ExportTransfer & KBOrders are template pages ready for client-specific business rules or removal.", cell_body_style)],
        [Paragraph("DO NOT Modify Without Review", cell_body_bold), Paragraph("src/components/modals/<br/>CreateDispatchModal.tsx<br/>src/types/index.ts", cell_body_bold), Paragraph("Vehicle payload calculation, stop sequencing, and PI allocation algorithms are tightly coupled with the dispatch workflow.", cell_body_style)],
        [Paragraph("Hub & Location Config", cell_body_bold), Paragraph("src/components/Navbar.tsx<br/>src/pages/DashboardPage.tsx", cell_body_bold), Paragraph("Hub name ('Bhiwandi Hub') is currently hardcoded in header badges. Can be made dynamic from user profile/auth context.", cell_body_style)]
    ]

    t_audit_pdf = Table(audit_rows, colWidths=[110, 140, 254])
    t_audit_pdf.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0F766E')),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
        ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('LEFTPADDING', (0, 0), (-1, -1), 5),
        ('RIGHTPADDING', (0, 0), (-1, -1), 5),
    ]))
    elements.append(t_audit_pdf)
    elements.append(Spacer(1, 10))

    # Section 6: Backend API Blueprint
    elements.append(Paragraph("6. Backend Team Integration Blueprint (API Contracts)", h1_style))
    elements.append(Paragraph("To connect the Frontend to a live Backend (Node.js / Spring Boot / Python / Go), implement these endpoints:", body_style))
    elements.append(Spacer(1, 6))

    api_rows = [
        [Paragraph("Method", cell_hdr_style), Paragraph("Endpoint Route", cell_hdr_style), Paragraph("Frontend Consumer Method", cell_hdr_style), Paragraph("Payload / Response Details", cell_hdr_style)],
        [Paragraph("GET", cell_body_bold), Paragraph("/api/v1/pis", cell_body_bold), Paragraph("dispatchService.getPIs()", cell_body_style), Paragraph("Returns list of Proforma Invoices", cell_body_style)],
        [Paragraph("POST", cell_body_bold), Paragraph("/api/v1/pis", cell_body_bold), Paragraph("dispatchService.addPI(piData)", cell_body_style), Paragraph("Creates new PI with line items", cell_body_style)],
        [Paragraph("PATCH", cell_body_bold), Paragraph("/api/v1/pis/:id/status", cell_body_bold), Paragraph("dispatchService.updatePIStatus(id, status)", cell_body_style), Paragraph("Approve, reject, or mark dispatched", cell_body_style)],
        [Paragraph("GET", cell_body_bold), Paragraph("/api/v1/dispatches", cell_body_bold), Paragraph("dispatchService.getDispatchPlans()", cell_body_style), Paragraph("Returns all dispatch plans with stops & vehicles", cell_body_style)],
        [Paragraph("POST", cell_body_bold), Paragraph("/api/v1/dispatches", cell_body_bold), Paragraph("dispatchService.createDispatchPlan(plan)", cell_body_style), Paragraph("Creates trip, reserves vehicle, marks PIs planned", cell_body_style)],
        [Paragraph("POST", cell_body_bold), Paragraph("/api/v1/dispatches/:id/rollback", cell_body_bold), Paragraph("dispatchService.rollbackDispatch(id, reason)", cell_body_style), Paragraph("Rolls back trip, releases vehicle, resets PIs to PENDING", cell_body_style)],
        [Paragraph("GET", cell_body_bold), Paragraph("/api/v1/fleet/vehicles", cell_body_bold), Paragraph("dispatchService.getVehicles()", cell_body_style), Paragraph("Returns vehicle fleet with capacity & availability", cell_body_style)],
        [Paragraph("GET", cell_body_bold), Paragraph("/api/v1/fleet/drivers", cell_body_bold), Paragraph("dispatchService.getDrivers()", cell_body_style), Paragraph("Returns driver master data", cell_body_style)],
        [Paragraph("GET", cell_body_bold), Paragraph("/api/v1/special-requests", cell_body_bold), Paragraph("dispatchService.getSpecialRequests()", cell_body_style), Paragraph("Returns supply chain special work requests", cell_body_style)],
        [Paragraph("POST", cell_body_bold), Paragraph("/api/v1/special-requests", cell_body_bold), Paragraph("dispatchService.addSpecialRequest(req)", cell_body_style), Paragraph("Creates new special request with attachments", cell_body_style)],
        [Paragraph("PATCH", cell_body_bold), Paragraph("/api/v1/special-requests/:id/done", cell_body_bold), Paragraph("dispatchService.updateSpecialRequestStatus(id, 'WORK_DONE')", cell_body_style), Paragraph("Marks request as work completed", cell_body_style)]
    ]

    t_api_pdf = Table(api_rows, colWidths=[45, 140, 155, 164])
    t_api_pdf.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1E293B')),
        ('ALIGN', (0, 0), (-1, -1), 'LEFT'),
        ('VALIGN', (0, 0), (-1, -1), 'TOP'),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E2E8F0')),
        ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
        ('TOPPADDING', (0, 0), (-1, -1), 3),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ('LEFTPADDING', (0, 0), (-1, -1), 5),
        ('RIGHTPADDING', (0, 0), (-1, -1), 5),
    ]))
    elements.append(t_api_pdf)

    doc.build(elements, canvasmaker=NumberedCanvas)
    print(f"[SUCCESS] PDF document generated: {filename}")


if __name__ == "__main__":
    docx_file = "Dispatch_Planning_Engineering_Handover_Document.docx"
    pdf_file = "Dispatch_Planning_Engineering_Handover_Document.pdf"
    create_word_document(docx_file)
    create_pdf_document(pdf_file)
    print("\n[ALL DONE] Handover documents (Word + PDF) generated successfully!")
